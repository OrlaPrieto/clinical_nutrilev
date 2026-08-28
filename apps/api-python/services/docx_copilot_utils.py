from datetime import datetime
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────────────────────────────────────────────────────────
# COLORES DEL TEMA NUTRILEV
# ─────────────────────────────────────────────────────────────────────────────
COLORS = {
    "primary_purple": "976CA0",
    "primary_teal":   "1A7A6E",
    "primary_rose":   "E8762B",
    "header_dark":     "1E293B",
    "header_light":    "F8FAFC",

    # Colores por tiempo de comida
    "desayuno_bg":     "FFFBEB", "desayuno_border": "F59E0B", "desayuno_text": "92400E",
    "col1_bg":         "ECFDF5", "col1_border":     "10B981", "col1_text": "065F46",
    "comida_bg":       "FFF1F2", "comida_border":   "F43F5E", "comida_text": "9F1239",
    "col2_bg":         "F5F3FF", "col2_border":     "8B5CF6", "col2_text": "5B21B6",
    "cena_bg":         "EEF2FF", "cena_border":     "6366F1", "cena_text": "3730A3",

    # Macros
    "protein_bg":      "EEF2FF", "protein_text":   "4338CA",
    "carbs_bg":        "FEF3C7", "carbs_text":     "B45309",
    "fat_bg":          "FFE4E6", "fat_text":       "BE123C",

    # Bordes y neutros
    "border_light":    "CBD5E1",
    "border_dark":     "94A3B8",
    "text_dark":       "0F172A",
    "text_mid":        "334155",
    "text_muted":      "64748B",
    "white":           "FFFFFF",
}

# ─────────────────────────────────────────────────────────────────────────────
# HELPERS DE FORMATO Y OXML
# ─────────────────────────────────────────────────────────────────────────────

def _rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

def _set_cell_bg(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.upper().lstrip("#"))
    tcPr.append(shd)

def _set_cell_borders(cell, color: str = "CBD5E1", size: int = 4):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color.upper().lstrip("#"))
        tcBorders.append(el)
    tcPr.append(tcBorders)

def _set_cell_margins(cell, top=40, left=60, bottom=40, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)

def _set_cell_valign(cell, align: str = "center"):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), align)
    tcPr.append(vAlign)

def _set_table_width(table, width_dxa: int = 14930):
    tblPr = table._tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(width_dxa))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

def _apply_col_widths(table, widths: list):
    tbl = table._tbl
    tblGrid = OxmlElement("w:tblGrid")
    for w in widths:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(w))
        tblGrid.append(gridCol)
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is not None:
        tblPr.addnext(tblGrid)
    else:
        tbl.insert(0, tblGrid)

def _add_run(para, text: str, bold=False, size_pt=8.0, color="0F172A", font="Arial", italic=False):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.color.rgb = _rgb(color)
    return run

def _format_date(date_val) -> str:
    if not date_val:
        return datetime.now().strftime("%d/%m/%Y")
    s = str(date_val)[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    return s

def _clean_ingredient(ing: str) -> str:
    if not ing:
        return ""
    import re
    return re.sub(r'^[•\-\*\s]+', '', str(ing)).strip()

def _set_cell_no_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "FFFFFF")
        tcBorders.append(el)
    tcPr.append(tcBorders)

def _add_paragraph_spacing(doc, before_pt=1, after_pt=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before_pt)
    p.paragraph_format.space_after = Pt(after_pt)

# ─────────────────────────────────────────────────────────────────────────────
# GENERADOR PRINCIPAL DOCX (LANDSCAPE 1 PÁGINA)
# ─────────────────────────────────────────────────────────────────────────────

def build_copilot_single_page_docx(copilot_data: dict, patient_context: dict, target_calories: int = 1800) -> bytes:
    doc = Document()

    # 1. Configurar Orientación LANDSCAPE y Márgenes Estrechos (0.8 cm)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(27.94)
    section.page_height = Cm(21.59)
    section.left_margin = Cm(0.8)
    section.right_margin = Cm(0.8)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)

    patient_name = patient_context.get("nombre") or patient_context.get("name") or "Paciente"
    fecha_hoy = _format_date(patient_context.get("fecha_hoy") or datetime.now())
    is_weekly = copilot_data.get("format_type") == "semanal" or bool(copilot_data.get("days"))

    # Extraer macros si existen
    clinical_analysis = copilot_data.get("clinical_analysis", {})
    macros = clinical_analysis.get("macro_distribution", {})
    cal_target = macros.get("calories", target_calories)
    prot_g = macros.get("protein_g", "—")
    prot_pct = macros.get("protein_pct", "")
    carb_g = macros.get("carbs_g", "—")
    carb_pct = macros.get("carbs_pct", "")
    fat_g = macros.get("fat_g", "—")
    fat_pct = macros.get("fat_pct", "")

    # 2. ENCABEZADO COMPACTO (Tabla de 1 fila, 2 columnas: Datos Paciente + Macros)
    head_tbl = doc.add_table(rows=1, cols=2)
    head_tbl.autofit = False
    _set_table_width(head_tbl, 14930)
    _apply_col_widths(head_tbl, [8930, 6000])

    left_cell = head_tbl.cell(0, 0)
    right_cell = head_tbl.cell(0, 1)

    _set_cell_bg(left_cell, "1E293B")
    _set_cell_bg(right_cell, "0F172A")
    _set_cell_margins(left_cell, top=60, left=120, bottom=60, right=120)
    _set_cell_margins(right_cell, top=60, left=120, bottom=60, right=120)
    _set_cell_no_borders(left_cell)
    _set_cell_no_borders(right_cell)

    # Info Paciente (Izquierda)
    p_left = left_cell.paragraphs[0]
    p_left.paragraph_format.space_before = Pt(0)
    p_left.paragraph_format.space_after = Pt(2)
    _add_run(p_left, "CLINICAL NUTRILIEV  ·  PLAN DE MENÚ COPILOTO\n", bold=True, size_pt=10.0, color="F8FAFC")
    _add_run(p_left, f"PACIENTE: {patient_name.upper()}   |   FECHA: {fecha_hoy}   |   OBJETIVO: {cal_target} kcal", bold=True, size_pt=8.0, color="CBD5E1")
    _add_run(p_left, f"   |   FORMATO: {'SEMANAL (7 DÍAS)' if is_weekly else 'EQUIVALENCIAS (3 OPCIONES)'}", italic=True, size_pt=7.5, color="94A3B8")

    # Macros Badges (Derecha)
    p_right = right_cell.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_before = Pt(0)
    p_right.paragraph_format.space_after = Pt(0)
    _add_run(p_right, f"Proteína: {prot_g}g ({prot_pct}%)   ", bold=True, size_pt=8.0, color="A5B4FC")
    _add_run(p_right, f"Carbohidratos: {carb_g}g ({carb_pct}%)   ", bold=True, size_pt=8.0, color="FDE68A")
    _add_run(p_right, f"Grasas: {fat_g}g ({fat_pct}%)", bold=True, size_pt=8.0, color="FECDD3")

    _add_paragraph_spacing(doc, 2, 2)

    # 3. CONSTRUCCIÓN DE LA TABLA PRINCIPAL DE MENÚS
    if is_weekly:
        _build_weekly_table(doc, copilot_data.get("days", []))
    else:
        _build_equivalencias_table(doc, copilot_data.get("menus", []))

    # 4. PIE DE PÁGINA CLÍNICO COMPACTO
    _add_paragraph_spacing(doc, 2, 2)
    footer_tbl = doc.add_table(rows=1, cols=1)
    footer_tbl.autofit = False
    _set_table_width(footer_tbl, 14930)
    f_cell = footer_tbl.cell(0, 0)
    _set_cell_bg(f_cell, "F8FAFC")
    _set_cell_borders(f_cell, "CBD5E1", 3)
    _set_cell_margins(f_cell, top=40, left=100, bottom=40, right=100)

    p_f = f_cell.paragraphs[0]
    p_f.paragraph_format.space_before = Pt(0)
    p_f.paragraph_format.space_after = Pt(0)
    
    rationale = clinical_analysis.get("clinical_rationale", "")
    if rationale:
        _add_run(p_f, "💡 ESTRATEGIA CLÍNICA: ", bold=True, size_pt=7.5, color="1A7A6E")
        _add_run(p_f, f"{rationale[:220]}... ", size_pt=7.0, color="334155")
        
    _add_run(p_f, "⚠️ Cualquier duda o incomodidad con los platillos, notifique a su nutrióloga. Clinical Nutrilev.", italic=True, size_pt=7.0, color="64748B")

    # Guardar a buffer binario
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_weekly_table(doc, days: list):
    """Construye tabla para 7 Días (Semanal) en 8 columnas (Tiempo + Lunes..Domingo)"""
    if not days:
        return

    col_widths = [2330] + [1800] * 7  # Total = 14930 dxa (~26.34 cm)
    meals_config = [
        ("desayuno", "🍳 DESAYUNO", COLORS["desayuno_bg"], COLORS["desayuno_border"], COLORS["desayuno_text"]),
        ("colacion_matutina", "🍏 COLACIÓN 1", COLORS["col1_bg"], COLORS["col1_border"], COLORS["col1_text"]),
        ("comida", "🍲 COMIDA", COLORS["comida_bg"], COLORS["comida_border"], COLORS["comida_text"]),
        ("colacion_vespertina", "🫐 COLACIÓN 2", COLORS["col2_bg"], COLORS["col2_border"], COLORS["col2_text"]),
        ("cena", "🌙 CENA", COLORS["cena_bg"], COLORS["cena_border"], COLORS["cena_text"]),
    ]

    # Filtrar solo tiempos que tengan al menos 1 platillo
    active_meals = []
    for key, label, bg, border, text_col in meals_config:
        if any(d.get(key) for d in days):
            active_meals.append((key, label, bg, border, text_col))

    table = doc.add_table(rows=1 + len(active_meals), cols=8)
    table.autofit = False
    _set_table_width(table, 14930)
    _apply_col_widths(table, col_widths)

    # 1. Row Header
    header_row = table.rows[0]
    cell_0 = header_row.cells[0]
    _set_cell_bg(cell_0, "1E293B")
    _set_cell_margins(cell_0, 40, 60, 40, 60)
    _set_cell_valign(cell_0, "center")
    p0 = cell_0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    _add_run(p0, "TIEMPO", bold=True, size_pt=8.0, color="FFFFFF")

    WEEKDAYS_ES = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO"]
    today_idx = datetime.now().weekday()
    start_idx = (today_idx + 1) % 7
    day_names = [WEEKDAYS_ES[(start_idx + i) % 7] for i in range(7)]
    for i in range(7):
        cell = header_row.cells[i + 1]
        d_name = days[i].get("day_name", day_names[i]).upper() if (i < len(days) and days[i].get("day_name")) else day_names[i]
        _set_cell_bg(cell, "0F172A")
        _set_cell_margins(cell, 40, 40, 40, 40)
        _set_cell_valign(cell, "center")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        _add_run(p, d_name, bold=True, size_pt=7.5, color="FFFFFF")

    # 2. Meal Rows
    for r_idx, (key, label, bg, border, text_col) in enumerate(active_meals):
        row = table.rows[r_idx + 1]
        
        # Col 0: Tiempo de comida
        t_cell = row.cells[0]
        _set_cell_bg(t_cell, bg)
        _set_cell_borders(t_cell, border, 4)
        _set_cell_margins(t_cell, 40, 60, 40, 60)
        _set_cell_valign(t_cell, "center")
        p_t = t_cell.paragraphs[0]
        p_t.paragraph_format.space_before = Pt(0)
        p_t.paragraph_format.space_after = Pt(0)
        _add_run(p_t, label, bold=True, size_pt=7.5, color=text_col)

        # Cols 1..7: Días
        for d_idx in range(7):
            cell = row.cells[d_idx + 1]
            _set_cell_borders(cell, "CBD5E1", 2)
            _set_cell_margins(cell, 30, 40, 30, 40)
            _set_cell_valign(cell, "top")
            
            day_data = days[d_idx] if d_idx < len(days) else {}
            dish = day_data.get(key)
            
            p_cell = cell.paragraphs[0]
            p_cell.paragraph_format.space_before = Pt(0)
            p_cell.paragraph_format.space_after = Pt(0)
            p_cell.paragraph_format.line_spacing = 1.0

            if dish and dish.get("platillo"):
                _add_run(p_cell, dish.get("platillo"), bold=True, size_pt=7.0, color="0F172A")
                ings = dish.get("ingredientes", [])
                if ings:
                    clean_ings = [_clean_ingredient(ing) for ing in ings if ing]
                    _add_run(p_cell, "\n" + ", ".join(clean_ings), size_pt=6.2, color="475569")
                if dish.get("preparacion_rapida"):
                    _add_run(p_cell, f"\n💡 {dish.get('preparacion_rapida')}", italic=True, size_pt=6.0, color="64748B")
            else:
                _add_run(p_cell, "—", italic=True, size_pt=6.5, color="94A3B8")


def _build_equivalencias_table(doc, menus: list):
    """
    Construye la tabla oficial de Menú de Equivalencias Nutrilev (SMAE):
    - Soporta dinámicamente 3 o 5 columnas de menú (Eq | Grupo | MENÚ 1 | MENÚ 2 | MENÚ 3 | [MENÚ 4] | [MENÚ 5])
    - Columnas Eq y Grupo compactas para maximizar el ancho de los platillos
    - Paleta sobria y profesional (Slate / Nutrilev) sin colores estridentes
    - Incluye todos los tiempos de comida (Desayuno, Colación Matutina, Comida, Colación Vespertina, Cena)
      con Grupo, Eq y opciones correspondientes a cada columna de menú.
    """
    if not menus:
        return

    num_menus = len(menus)
    num_cols = 2 + num_menus

    # Anchos calculados: Eq y Grupo compactos
    # Para 3 menús: Eq 480, Grupo 980, Menús 4490 c/u (Total 14930 dxa)
    # Para 5 menús: Eq 450, Grupo 900, Menús 2716 c/u (Total 14930 dxa)
    eq_w = 450 if num_menus >= 5 else 480
    grupo_w = 900 if num_menus >= 5 else 980
    rem_w = 14930 - (eq_w + grupo_w)
    each_menu_w = int(rem_w / num_menus)
    col_widths = [eq_w, grupo_w] + [each_menu_w] * (num_menus - 1) + [rem_w - each_menu_w * (num_menus - 1)]

    # Tamaño de tipografía dinámico según cantidad de columnas
    dish_font_size = 6.2 if num_menus >= 5 else 7.0
    tip_font_size = 5.8 if num_menus >= 5 else 6.5

    all_meals = [
        ("desayuno", "🍳 DESAYUNO"),
        ("colacion_matutina", "🍏 COLACIÓN MATUTINA"),
        ("comida", "🍲 COMIDA"),
        ("colacion_vespertina", "🫐 COLACIÓN VESPERTINA"),
        ("cena", "🌙 CENA")
    ]

    rows_to_render = []

    for m_key, m_label in all_meals:
        # Verificar si algún menú tiene este tiempo
        if not any(m.get(m_key) for m in menus):
            continue

        # 1. Section Header Row
        rows_to_render.append(("section_header", m_label))

        # 2. Column Titles Row (Eq, Grupo, MENÚ 1..N)
        dish_titles = []
        for idx, m in enumerate(menus):
            meal_obj = m.get(m_key, {})
            if isinstance(meal_obj, dict):
                t = meal_obj.get("nombre_platillo") or meal_obj.get("platillo") or meal_obj.get("title") or f"MENÚ {idx + 1}"
            else:
                t = f"MENÚ {idx + 1}"
            dish_titles.append(t)

        rows_to_render.append(("col_header", dish_titles))

        # 3. Extraer grupos presentes en este tiempo de comida
        groups_in_meal = []
        for m in menus:
            meal_data = m.get(m_key, {})
            if isinstance(meal_data, dict):
                eqs = meal_data.get("equivalencias", [])
                if isinstance(eqs, list):
                    for eq in eqs:
                        if isinstance(eq, dict):
                            g = eq.get("grupo")
                            if g and g not in groups_in_meal:
                                groups_in_meal.append(g)

        # Fallback de grupos estándar según el tipo de comida
        if not groups_in_meal:
            if "colacion" in m_key:
                groups_in_meal = ["Frutas", "Verduras"]
            else:
                groups_in_meal = ["Cereales", "POA", "Grasas", "Frutas", "Verduras"]

        # Ordenar grupos en secuencia oficial SMAE
        standard_order = ["Cereales", "POA", "Grasas", "Frutas", "Verduras"]
        groups_in_meal.sort(key=lambda x: standard_order.index(x) if x in standard_order else 99)

        # 4. Construir filas de grupos para este tiempo
        for g_name in groups_in_meal:
            eq_val = ""
            m_descs = []

            for idx, m in enumerate(menus):
                meal_data = m.get(m_key, {})
                desc = "—"

                if isinstance(meal_data, dict):
                    eqs = meal_data.get("equivalencias", [])
                    if isinstance(eqs, list) and eqs:
                        found_eq = next((eq for eq in eqs if isinstance(eq, dict) and eq.get("grupo", "").lower() == g_name.lower()), None)
                        if found_eq:
                            if not eq_val and found_eq.get("porciones"):
                                eq_val = str(found_eq.get("porciones"))
                            desc = found_eq.get("descripcion") or "—"
                    elif meal_data.get("descripcion"):
                        desc = str(meal_data.get("descripcion"))
                        if not eq_val:
                            eq_val = "1"
                    elif meal_data.get("ingredientes"):
                        clean_ings = [_clean_ingredient(i) for i in meal_data.get("ingredientes", []) if i]
                        if clean_ings:
                            desc = ", ".join(clean_ings)
                        if not eq_val:
                            eq_val = "1"

                m_descs.append(desc)

            if not eq_val:
                eq_val = "1"

            rows_to_render.append(("group_row", (eq_val, g_name, m_descs)))

    # Crear tabla unificada en Word
    table = doc.add_table(rows=len(rows_to_render), cols=num_cols)
    table.autofit = False
    _set_table_width(table, 14930)
    _apply_col_widths(table, col_widths)

    for r_idx, (r_type, r_data) in enumerate(rows_to_render):
        row = table.rows[r_idx]

        if r_type == "section_header":
            m_label = r_data
            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                _set_cell_bg(cell, "1E293B")
                _set_cell_no_borders(cell)
                _set_cell_margins(cell, top=25, left=50, bottom=25, right=50)
                _set_cell_valign(cell, "center")
                if c_idx == 0:
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    _add_run(p, f"  {m_label}", bold=True, size_pt=8.0, color="FFFFFF")

        elif r_type == "col_header":
            dish_titles = r_data
            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                _set_cell_margins(cell, top=25, left=35, bottom=25, right=35)
                _set_cell_valign(cell, "center")
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)

                if c_idx == 0:
                    _set_cell_bg(cell, "1E293B")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_run(p, "Eq", bold=True, size_pt=7.2, color="FFFFFF")
                elif c_idx == 1:
                    _set_cell_bg(cell, "334155")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_run(p, "Grupo", bold=True, size_pt=7.2, color="FFFFFF")
                else:
                    opt_idx = c_idx - 2
                    title_text = dish_titles[opt_idx] if opt_idx < len(dish_titles) else f"MENÚ {opt_idx + 1}"
                    _set_cell_bg(cell, "334155")
                    _set_cell_borders(cell, "475569", 2)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_run(p, f"MENÚ {opt_idx + 1}: {title_text.upper()}", bold=True, size_pt=dish_font_size, color="FFFFFF")

        elif r_type == "group_row":
            eq_val, g_name, m_descs = r_data

            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                _set_cell_borders(cell, "E2E8F0", 2)
                _set_cell_margins(cell, 20, 35, 20, 35)
                _set_cell_valign(cell, "center")
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0

                if c_idx == 0:
                    _set_cell_bg(cell, "F1F5F9")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_run(p, str(eq_val), bold=True, size_pt=7.2, color="1E293B")
                elif c_idx == 1:
                    _set_cell_bg(cell, "F8FAFC")
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    _add_run(p, g_name, bold=True, size_pt=7.0, color="334155")
                else:
                    opt_idx = c_idx - 2
                    desc_text = m_descs[opt_idx] if opt_idx < len(m_descs) else "—"
                    _set_cell_bg(cell, "FFFFFF" if r_idx % 2 == 0 else "FAFAFA")
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    _add_run(p, desc_text, size_pt=dish_font_size, color="0F172A")

